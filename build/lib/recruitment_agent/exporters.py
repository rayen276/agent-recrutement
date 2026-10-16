"""Exports : rapport Markdown, lettre et rapport Word, JSON."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .models import Analysis, InterviewPrep, Result

_STATUT = {"couvert": "✅ Couvert", "partiel": "🟡 Partiel", "absent": "❌ Absent"}


def report_markdown(a: Analysis) -> str:
    lines = [
        f"# Analyse de candidature — {a.poste}" + (f" ({a.entreprise})" if a.entreprise else ""),
        "",
        f"**Score de correspondance : {a.score}/100**",
        "",
        a.resume,
        "",
        "## Points forts",
        *[f"- **{s.titre}** — {s.preuve}" for s in a.points_forts],
        "",
        "## Écarts et conseils",
        *([f"- **{g.titre}** — {g.detail}\n  - Conseil : {g.conseil}" for g in a.ecarts] or ["- Aucun écart majeur."]),
        "",
        "## Exigences de l'offre",
        "| Exigence | Statut | Commentaire |",
        "|---|---|---|",
        *[f"| {r.exigence} | {_STATUT[r.statut]} | {r.commentaire} |" for r in a.exigences],
        "",
        "## Mots-clés",
        f"- Présents : {', '.join(a.mots_cles.presents) or '—'}",
        f"- Manquants : {', '.join(a.mots_cles.manquants) or '—'}",
        "",
        "## Pour améliorer le CV",
        *[f"- {c}" for c in a.conseils_cv],
        "",
    ]
    return "\n".join(lines)


def letter_docx_bytes(letter: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for block in [b.strip() for b in letter.split("\n\n") if b.strip()]:
        doc.add_paragraph(block)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def report_docx_bytes(a: Analysis) -> bytes:
    doc = Document()
    doc.add_heading(f"Analyse de candidature — {a.poste}", level=1)
    if a.entreprise:
        doc.add_paragraph(a.entreprise)
    p = doc.add_paragraph()
    p.add_run(f"Score de correspondance : {a.score}/100").bold = True
    doc.add_paragraph(a.resume)

    doc.add_heading("Points forts", level=2)
    for s in a.points_forts:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(s.titre).bold = True
        para.add_run(f" — {s.preuve}")

    doc.add_heading("Écarts et conseils", level=2)
    for g in a.ecarts:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(g.titre).bold = True
        para.add_run(f" — {g.detail} Conseil : {g.conseil}")
    if not a.ecarts:
        doc.add_paragraph("Aucun écart majeur.")

    doc.add_heading("Exigences de l'offre", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, title in zip(table.rows[0].cells, ("Exigence", "Statut", "Commentaire")):
        cell.text = title
    for r in a.exigences:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = r.exigence, r.statut.capitalize(), r.commentaire

    doc.add_heading("Mots-clés", level=2)
    doc.add_paragraph(f"Présents : {', '.join(a.mots_cles.presents) or '—'}")
    doc.add_paragraph(f"Manquants : {', '.join(a.mots_cles.manquants) or '—'}")

    doc.add_heading("Pour améliorer le CV", level=2)
    for c in a.conseils_cv:
        doc.add_paragraph(c, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def save_all(result: Result, out_dir: str | Path) -> dict[str, Path]:
    """Écrit tous les livrables dans un dossier et renvoie leurs chemins."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    files = {
        "analyse.json": result.analysis.model_dump_json(indent=2).encode("utf-8"),
        "rapport.md": report_markdown(result.analysis).encode("utf-8"),
        "rapport.docx": report_docx_bytes(result.analysis),
        "lettre.md": (result.letter + "\n").encode("utf-8"),
        "lettre.docx": letter_docx_bytes(result.letter),
    }
    paths: dict[str, Path] = {}
    for name, data in files.items():
        path = out / name
        path.write_bytes(data)
        paths[name] = path
    return paths


def interview_markdown(prep: InterviewPrep) -> str:
    lines = ["# Préparation d'entretien", ""]
    for i, q in enumerate(prep.questions, 1):
        lines += [
            f"## {i}. {q.question}",
            f"*Catégorie : {q.categorie}*",
            "",
            f"**Ce que le recruteur évalue :** {q.objectif}",
            "",
            f"**Piste de réponse :** {q.piste_reponse}",
            "",
        ]
    lines += ["## Questions à poser au recruteur", *[f"- {x}" for x in prep.questions_a_poser], ""]
    lines += ["## Conseils", *[f"- {x}" for x in prep.conseils], ""]
    return "\n".join(lines)


def interview_docx_bytes(prep: InterviewPrep) -> bytes:
    doc = Document()
    doc.add_heading("Préparation d'entretien", level=1)
    for i, q in enumerate(prep.questions, 1):
        doc.add_heading(f"{i}. {q.question}", level=3)
        doc.add_paragraph(f"Catégorie : {q.categorie}").italic = True
        p = doc.add_paragraph()
        p.add_run("Ce que le recruteur évalue : ").bold = True
        p.add_run(q.objectif)
        p = doc.add_paragraph()
        p.add_run("Piste de réponse : ").bold = True
        p.add_run(q.piste_reponse)
    doc.add_heading("Questions à poser au recruteur", level=2)
    for x in prep.questions_a_poser:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("Conseils", level=2)
    for x in prep.conseils:
        doc.add_paragraph(x, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_zip(analysis: Analysis | None, letter: str = "", prep: InterviewPrep | None = None) -> bytes:
    """Dossier de candidature complet (rapport, lettre, fiche d'entretien) dans une archive ZIP."""
    import zipfile

    if analysis is None and not letter and prep is None:
        return b""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        if analysis is not None:
            z.writestr("rapport.md", report_markdown(analysis))
            z.writestr("rapport.docx", report_docx_bytes(analysis))
        if letter.strip():
            z.writestr("lettre.md", letter + "\n")
            z.writestr("lettre.docx", letter_docx_bytes(letter))
        if prep is not None:
            z.writestr("entretien.md", interview_markdown(prep))
            z.writestr("entretien.docx", interview_docx_bytes(prep))
    return buf.getvalue()
