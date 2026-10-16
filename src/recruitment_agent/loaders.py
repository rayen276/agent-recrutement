"""Lecture des CV et offres : .txt, .md, .pdf, .docx."""

from __future__ import annotations

import io
from pathlib import Path
from typing import BinaryIO

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


class LoaderError(ValueError):
    """Fichier illisible ou format non supporté."""


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(stream: BinaryIO) -> str:
    from pypdf import PdfReader

    reader = PdfReader(stream)
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    if not text:
        raise LoaderError(
            "Aucun texte extrait du PDF (document scanné ?). "
            "Utilisez un PDF texte ou collez le contenu dans un fichier .txt."
        )
    return text


def _read_docx(stream: BinaryIO) -> str:
    from docx import Document

    doc = Document(stream)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_bytes(filename: str, data: bytes) -> str:
    """Extrait le texte d'un fichier fourni en mémoire (ex. upload web)."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise LoaderError(f"Format non supporté : {ext or 'inconnu'}. Formats : {', '.join(sorted(SUPPORTED))}")
    try:
        if ext in {".txt", ".md"}:
            text = _decode(data)
        elif ext == ".pdf":
            text = _read_pdf(io.BytesIO(data))
        else:
            text = _read_docx(io.BytesIO(data))
    except LoaderError:
        raise
    except Exception as exc:  # fichier corrompu, etc.
        raise LoaderError(f"Impossible de lire « {filename} » : {exc}") from exc
    return clean_text(text)


def load_file(path: str | Path) -> str:
    p = Path(path)
    if not p.is_file():
        raise LoaderError(f"Fichier introuvable : {p}")
    return load_bytes(p.name, p.read_bytes())


def clean_text(text: str) -> str:
    """Normalise les sauts de ligne et supprime les espaces superflus."""
    lines = [ln.rstrip() for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if ln.strip():
            blank = 0
            out.append(ln)
        else:
            blank += 1
            if blank == 1:
                out.append("")
    return "\n".join(out).strip()
